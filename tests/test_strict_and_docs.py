"""Tests for the --strict trust gate and the docs/dependency surfaces.

The trust gate is the contract between the exploratory matcher and any
real-file pilot: it must fail loudly when an eligible Word number is
UNRESOLVED or LOW, and only then. The docs/dep tests guard the
reproducibility surface — README quickstart and requirements.txt — so a
fresh clone can be set up and run the same way every time.
"""
from __future__ import annotations

from pathlib import Path

import docx
import openpyxl
import pytest

from src.main import main as cli_main
from src.synthetic_generator import generate
from src.validator import CoverageSummary


REPO_ROOT = Path(__file__).resolve().parents[1]


# ---------------------------------------------------------------------------
# Strict gate — unit-level
# ---------------------------------------------------------------------------

def _summary(**counts: int) -> CoverageSummary:
    by = {"HIGH": 0, "MEDIUM": 0, "LOW": 0, "UNRESOLVED": 0, "EXCLUDED": 0}
    by.update(counts)
    return CoverageSummary(total=sum(by.values()), by_confidence=by)


def test_strict_failures_counts_unresolved_and_low_only():
    assert _summary(HIGH=2, MEDIUM=1, EXCLUDED=5).strict_failures == 0
    assert _summary(HIGH=1, LOW=2).strict_failures == 2
    assert _summary(UNRESOLVED=3, EXCLUDED=4).strict_failures == 3
    # EXCLUDED must never count against the gate — those are policy-skipped
    # with an audit-trail entry, not unsafe.
    assert _summary(EXCLUDED=10).strict_failures == 0


# ---------------------------------------------------------------------------
# Strict gate — CLI behaviour
# ---------------------------------------------------------------------------

def test_strict_fails_loudly_on_synthetic_corpus(tmp_path: Path, capsys):
    """The synthetic corpus contains deliberate UNRESOLVED cases (e.g. '15%',
    '0.70个百分点'). --strict must surface that as a non-zero exit and an
    unmistakable stderr message; the artifacts must still be written so a
    reviewer can inspect what failed."""
    samples = tmp_path / "samples"
    output = tmp_path / "output"
    generate(samples)

    rc = cli_main([
        "learn",
        "--excel", str(samples / "historical.xlsx"),
        "--word", str(samples / "finished_report.docx"),
        "--out", str(output),
        "--strict",
    ])

    assert rc == 3, "strict mode must return a non-zero exit code on UNRESOLVED/LOW"
    err = capsys.readouterr().err
    assert "STRICT GATE FAILED" in err
    assert "UNRESOLVED" in err
    # Artifacts must still exist — a failed gate doesn't mean a failed pipeline.
    assert (output / "mapping_review.xlsx").exists()
    assert (output / "confidence_report.md").exists()


def test_default_mode_warns_loudly_but_exits_zero(tmp_path: Path, capsys):
    """Without --strict, the same corpus must still write artifacts AND print
    an unmistakable warning to stderr — exploratory mode is not silent mode."""
    samples = tmp_path / "samples"
    output = tmp_path / "output"
    generate(samples)

    rc = cli_main([
        "learn",
        "--excel", str(samples / "historical.xlsx"),
        "--word", str(samples / "finished_report.docx"),
        "--out", str(output),
    ])

    assert rc == 0
    err = capsys.readouterr().err
    assert "WARNING" in err
    assert "UNRESOLVED" in err or "LOW" in err
    assert "--strict" in err, "warning must point the reader at the trust gate"
    assert (output / "mapping_review.xlsx").exists()
    assert (output / "confidence_report.md").exists()


def test_strict_passes_when_every_eligible_number_is_mapped(tmp_path: Path, capsys):
    """If the eligible set has no UNRESOLVED/LOW, --strict must exit 0 and
    emit no STRICT GATE FAILED banner. We build a minimal clean pair here so
    the test does not depend on the synthetic corpus (which is deliberately
    dirty)."""
    sample_dir = tmp_path / "samples"
    output = tmp_path / "output"
    sample_dir.mkdir()

    xlsx = sample_dir / "clean.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "月报"
    ws.append(["指标", "2026年5月"])
    ws.append(["营业收入", 100])
    wb.save(str(xlsx))

    docx_path = sample_dir / "clean.docx"
    doc = docx.Document()
    # "5" inside "5月" is captured as EXCLUDED (date marker) and does not
    # count against the gate; "100" with strong context → HIGH.
    doc.add_paragraph("5月营业收入达100元。")
    doc.save(str(docx_path))

    rc = cli_main([
        "learn",
        "--excel", str(xlsx),
        "--word", str(docx_path),
        "--out", str(output),
        "--strict",
    ])

    out_streams = capsys.readouterr()
    assert rc == 0, f"strict mode must pass on a clean pair; stderr was:\n{out_streams.err}"
    assert "STRICT GATE FAILED" not in out_streams.err
    assert "WARNING" not in out_streams.err
    assert (output / "mapping_review.xlsx").exists()
    assert (output / "confidence_report.md").exists()


# ---------------------------------------------------------------------------
# Docs + dependency surfaces
# ---------------------------------------------------------------------------

def test_requirements_lists_current_runtime_and_test_deps():
    """A fresh clone must be able to `pip install -r requirements.txt` and
    get exactly the packages this prototype imports — no more, no less.
    Adding a new top-level import without bumping this file is a bug."""
    req_path = REPO_ROOT / "requirements.txt"
    assert req_path.exists(), "requirements.txt must live at repo root"
    text = req_path.read_text(encoding="utf-8")

    for pkg in ("openpyxl", "python-docx", "PyYAML", "pytest"):
        assert pkg in text, f"requirements.txt missing {pkg}"

    # Guard against silent scope creep — anything beyond the deterministic
    # learn-mode quartet should be a deliberate, reviewed addition.
    forbidden = ("openai", "anthropic", "torch", "tensorflow", "langchain")
    lowered = text.lower()
    for pkg in forbidden:
        assert pkg not in lowered, f"requirements.txt unexpectedly lists {pkg}"


@pytest.mark.parametrize("phrase", [
    "pip install -r requirements.txt",
    "generate-synthetic",
    "mapping_review.xlsx",
    "confidence_report.md",
    "auto_mapping.yml",
    "converted_template.docx",
    "{{ word_NNNN }}",
    "HIGH",
    "MEDIUM",
    "LOW",
    "UNRESOLVED",
    "EXCLUDED",
    "--strict",
])
def test_readme_documents_quickstart_statuses_and_strict_gate(phrase: str):
    """The README is the only on-disk document a new contributor reads
    before touching the tool. It must cover setup, the learn artifacts, the
    five confidence statuses (including EXCLUDED audit rows), and the
    --strict trust gate. If any of these drop out, the surface that protects
    against silently misuing the prototype is gone."""
    readme = (REPO_ROOT / "README.md").read_text(encoding="utf-8")
    assert phrase in readme, f"README is missing required phrase: {phrase!r}"
